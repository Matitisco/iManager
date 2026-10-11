import json, random, openpyxl
from openpyxl.styles import Font
random.seed(7)
import os; OUT=os.path.dirname(os.path.abspath(__file__))
MODELS=[("11",["64","128"]),("12",["64","128"]),("12 Pro",["128","256"]),("13",["128","256"]),("13 mini",["128"]),("13 Pro Max",["256"]),("14",["128"]),("14 Pro",["128","256"]),("15",["128","256"]),("15 Pro Max",["256","512"]),("16",["128"]),("16 Pro",["256"]),("17 Pro Max",["256","512"])]
COLORS=["Black","White","Blue","Midnight","Starlight","Natural","Desert","Red","Purple"]
def dev(cond="USADO"):
    m,caps=random.choice(MODELS); c=random.choice(caps)
    bat=random.randint(78,100) if cond=="USADO" else None
    usd=random.randrange(250,1800,10)
    imei="35"+"".join(random.choice("0123456789") for _ in range(13))
    return dict(model=m,capacity=c+" GB",battery=bat,color=random.choice(COLORS),usd=usd,imei=imei,condition=cond)
def gt(d,price,imei=True,color=True,battery=True):
    r={"model":d["model"],"capacity":d["capacity"],"condition":d["condition"],"price":price}
    if color: r["color"]=d["color"]
    if imei: r["imei"]=d["imei"]
    if battery and d["battery"] is not None: r["batteryHealth"]=f'{d["battery"]}%'
    return r
meta={}
def save(wb,name,cur,rows,notes):
    wb.save(f"{OUT}/{name}.xlsx"); json.dump({"currency":cur,"rows":rows,"notes":notes},open(f"{OUT}/{name}.expected.json","w"),ensure_ascii=False,indent=1)

# 1 encabezado corrido + título combinado + filas vacías + baterías mezcladas
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Usados Premium"
ws.merge_cells("A1:F3"); ws["A1"]="iSTORE PALERMO — LISTA DE PRECIOS EQUIPOS USADOS"; ws["A1"].font=Font(bold=True,size=16)
ws["A5"]="Actualizado 10/10 - consultar disponibilidad por WhatsApp"
ws.append([]); ws.append(["","Equipo","Bat.","Color","Precio u$s","IMEI"])
exp=[]
for i in range(22):
    d=dev(); bf=random.choice(["%","frac","int"])
    b={"%":f'{d["battery"]}%',"frac":d["battery"]/100,"int":d["battery"]}[bf]
    ws.append(["",f'iPhone {d["model"]} {d["capacity"].replace(" ","")}',b,d["color"],d["usd"],d["imei"]])
    if i in (7,15): ws.append([]); ws.append([])
    exp.append(gt(d,d["usd"]) | {"model":"iPhone "+d["model"]})
ws.append([]); ws.append(["","Precios en efectivo. Transferencia +5%"])
save(wb,"f1_encabezado_corrido","USD",exp,"encabezado en fila 7 corrido a la columna B, título combinado A1:F3, filas vacías intercaladas, batería 87% / 0.87 / 87")

# 2 encabezado en dos filas (Precio combinado sobre USD / ARS)
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Lista"
ws["A1"]="MUNDO APPLE CÓRDOBA"; ws["A2"]="Dólar blue hoy: $1400"
ws.append([]); ws.append(["Modelo","Capacidad","Color","Precio",None,"Estado"]); ws.merge_cells("D4:E4")
for c in "ABCF": ws.merge_cells(f"{c}4:{c}5")
ws.append([None,None,None,"USD","ARS",None])
exp=[]
for i in range(18):
    d=dev(random.choice(["NUEVO","USADO"])); ws.append([d["model"],d["capacity"],d["color"],d["usd"],f'$ {d["usd"]*1400:,}'.replace(",","."),"Sellado" if d["condition"]=="NUEVO" else "Usado"])
    exp.append(gt(d,d["usd"],imei=False,battery=False))
save(wb,"f2_encabezado_dos_filas","USD",exp,"encabezado en dos filas con celdas combinadas (Precio sobre USD|ARS), estado en columna")

# 3 varias tablas en una hoja
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Stock"
ws.append(["STOCK OCTUBRE"]); ws.append([])
exp=[]
for title,cond in (("NUEVOS SELLADOS","NUEVO"),("USADOS / SEMINUEVOS","USADO")):
    ws.append([title]); ws.append(["Modelo","GB","Color","Batería","Precio USD","IMEI"])
    for i in range(12):
        d=dev(cond); ws.append([d["model"],int(d["capacity"].split()[0]),d["color"],f'{d["battery"]}%' if d["battery"] else "-",d["usd"],d["imei"]])
        exp.append(gt(d,d["usd"]))
    ws.append([]); ws.append([])
save(wb,"f3_varias_tablas","USD",exp,"dos tablas (nuevos y usados) en la misma hoja con encabezado repetido; condición solo en el título de cada tabla; capacidad como número")

# 4 ARS y USD mezclados en la misma columna, tienda en ARS
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Precios"
ws["A1"]="TECNO ROSARIO"; ws["D1"]="Cotización:"; ws["E1"]="$ 1.400"
ws.append([]); ws.append(["Mod","Alm","Bat","Precio","Obs"])
exp=[]
for i in range(20):
    d=dev(); k=i%3
    p={0:f'$ {d["usd"]*1400:,}'.replace(",","."),1:f'u$s {d["usd"]}',2:f'{d["usd"]} dls'}[k]
    ws.append([d["model"],d["capacity"].replace(" ",""),f'Bat {d["battery"]}%',p,random.choice(["","Detalle en marco","Caja original",""])])
    exp.append(gt(d,d["usd"]*1400,imei=False,color=False) )
save(wb,"f4_monedas_mezcladas","ARS",exp,"precio en una sola columna con '$ 1.200.000', 'u$s 850', '850 dls'; tienda en ARS, cotización en E1; capacidad pegada '128GB'; batería 'Bat 87%'")

# 5 hojas irrelevantes + columnas en otro orden
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Resumen"
ws.append(["Resumen mensual"]); ws.append(["Mes","Ventas","Ganancia"]); [ws.append([m,random.randint(20,60),random.randint(3000,9000)]) for m in ["Jul","Ago","Sep"]]
g=wb.create_sheet("Gráficos"); g.append(["Modelo","Unidades vendidas"]); [g.append([m,random.randint(1,20)]) for m,_ in MODELS[:6]]
s=wb.create_sheet("Stock Seminuevos"); s.append(["Precio venta USD","Costo","IMEI","Modelo","Capacidad","Color","Salud bat."])
exp=[]
for i in range(25):
    d=dev(); s.append([d["usd"],d["usd"]-random.randint(60,150),d["imei"],"iPhone "+d["model"],d["capacity"],d["color"],d["battery"]/100])
    e=gt(d,d["usd"]) | {"model":"iPhone "+d["model"]}; exp.append(e)
p=wb.create_sheet("Proveedores"); p.append(["Proveedor","Teléfono","Ciudad"]); p.append(["Distribuidora Sur","11 5555-1234","CABA"])
save(wb,"f5_hojas_irrelevantes","USD",exp,"hojas Resumen/Gráficos/Proveedores irrelevantes; columnas reordenadas, costo antes del modelo; batería 0.87")

# 6 hoja con miles de filas
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Inventario"
ws.append(["Estado","Modelo","Capacidad","Color","Batería","IMEI","Precio USD"])
exp=[]
for i in range(3000):
    d=dev(random.choice(["NUEVO","USADO","USADO"]))
    ws.append(["Nuevo" if d["condition"]=="NUEVO" else "Usado",d["model"],d["capacity"],d["color"],f'{d["battery"]}%' if d["battery"] else "",d["imei"],d["usd"]])
    exp.append(gt(d,d["usd"]))
    if i in (999,1999): ws.append([]); ws.append([])
save(wb,"f6_miles_de_filas","USD",exp,"3000 equipos, condición por columna, filas vacías en el medio")

# 7 nombres raros/abreviados
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Hoja1"
ws.append(["", "", "LISTA MAYORISTA - NO DIFUNDIR"]); ws.append([])
ws.append(["Descripcion","Cond","Pila","Tono","U$D","Nro serie/IMEI","Obs"])
exp=[]
for i in range(20):
    d=dev(random.choice(["NUEVO","USADO"]))
    ws.append([f'IPH {d["model"].upper()} {d["capacity"]}', "Sellado" if d["condition"]=="NUEVO" else "Usado", "" if d["battery"] is None else f'{d["battery"]} %', d["color"], d["usd"], d["imei"], random.choice(["","Reservado?",""])])
    exp.append(gt(d,d["usd"]) | {"model":"IPH "+d["model"].upper()})
ws.append(["TOTAL","","","",sum(e["price"] for e in exp)])
save(wb,"f7_nombres_raros","USD",exp,"encabezados abreviados (Pila, Tono, U$D, Cond), fila TOTAL al final")
print("ok")
